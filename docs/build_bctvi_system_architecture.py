from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parent
SOURCE_IMAGE_PATH = ROOT / "diagrams" / "CCTN-System-Architecture-SSC-style.png"
IMAGE_PATH = ROOT / "diagrams" / "BCTVI-System-Architecture-Detailed.png"
OUTPUT_PATH = ROOT / "BCTVI_SYSTEM_ARCHITECTURE_DETAILED.docx"


def set_cell_borderless_picture_alt_text(inline_shape, title: str, description: str) -> None:
    doc_properties = inline_shape._inline.docPr
    doc_properties.set("name", title)
    doc_properties.set("descr", description)


def font(name: str, size: int):
    return ImageFont.truetype(Path("C:/Windows/Fonts") / name, size)


diagram = Image.open(SOURCE_IMAGE_PATH).convert("RGB")
draw = ImageDraw.Draw(diagram)

# Update the system name while preserving the source diagram's layout.
draw.rectangle((1990, 435, 2825, 565), fill=(236, 235, 248))
draw.text(
    (2408, 455),
    "BCTVI Online Cable Service Management",
    anchor="ma",
    fill=(46, 43, 107),
    font=font("arialbd.ttf", 37),
)
draw.text(
    (2408, 502),
    "Laravel 9 | PHP 8 | Blade | Bootstrap | Sanctum",
    anchor="ma",
    fill=(95, 91, 158),
    font=font("arial.ttf", 24),
)
draw.text(
    (2408, 537),
    "Android client: Kotlin | Jetpack Compose",
    anchor="ma",
    fill=(95, 91, 158),
    font=font("arial.ttf", 22),
)

# The disconnection workflow is no longer part of the client portal.
draw.rectangle((795, 1190, 1685, 1320), fill=(255, 255, 255))
draw.line((795, 1190, 1685, 1190), fill=(226, 232, 240), width=2)
draw.line((795, 1320, 1685, 1320), fill=(226, 232, 240), width=2)
draw.text(
    (808, 1237),
    "Settings and plan management",
    fill=(47, 58, 72),
    font=font("arial.ttf", 29),
)
diagram.save(IMAGE_PATH, quality=95)

document = Document()
section = document.sections[0]
section.start_type = WD_SECTION_START.NEW_PAGE
section.page_width = Inches(14)
section.page_height = Inches(8.5)
section.top_margin = Inches(0.15)
section.bottom_margin = Inches(0.15)
section.left_margin = Inches(0.25)
section.right_margin = Inches(0.25)
section.header_distance = Inches(0.1)
section.footer_distance = Inches(0.1)

normal_style = document.styles["Normal"]
normal_style.font.name = "Arial"
normal_style.font.size = Pt(10)
normal_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")

title = document.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(1)
title_run = title.add_run("BCTVI ONLINE CABLE SERVICE MANAGEMENT SYSTEM")
title_run.bold = True
title_run.font.name = "Arial"
title_run.font.size = Pt(11)
title_run.font.color.rgb = RGBColor(0, 0, 0)

diagram_paragraph = document.add_paragraph()
diagram_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
diagram_paragraph.paragraph_format.space_before = Pt(0)
diagram_paragraph.paragraph_format.space_after = Pt(0)
diagram_run = diagram_paragraph.add_run()
shape = diagram_run.add_picture(str(IMAGE_PATH), width=Inches(13.35))
set_cell_borderless_picture_alt_text(
    shape,
    "BCTVI Detailed System Architecture",
    "Detailed multi-column BCTVI architecture showing users and portals, system modules, the application core, data storage, and external services.",
)

document.core_properties.title = "BCTVI System Architecture"
document.core_properties.subject = "Detailed multi-column system architecture diagram"
document.core_properties.author = "BCTVI Capstone Project Team"
document.core_properties.keywords = "BCTVI, system architecture, Laravel, Android, MySQL"

document.save(OUTPUT_PATH)
print(OUTPUT_PATH)
